"""openpyxl 只读 Excel 工具：sheet 概览 + 表头/示例行/按范围读取。"""
from datetime import date, datetime
from pathlib import Path
import re

from openpyxl import load_workbook


EXCEL_ERROR_VALUES = {"#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}


def _cell(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        cleaned = value.strip()
        if not cleaned or cleaned.upper() in EXCEL_ERROR_VALUES:
            return None
        return cleaned
    return value


def _col_number(part: str) -> int:
    result = 0
    for ch in part:
        if ch.isalpha():
            result = result * 26 + (ord(ch.upper()) - ord("A") + 1)
    return result or 1


def _row_number(part: str) -> int | None:
    digits = "".join(ch for ch in part if ch.isdigit())
    return int(digits) if digits else None


def _parse_range(range_text: str, max_row: int, max_col: int) -> tuple[int, int, int, int]:
    """把 A2:D15 或 Sheet1!D10:A1 或 A1 解析为纠偏后的 (start_row, start_col, end_row, end_col)。"""
    if not range_text:
        return 1, 1, max(1, max_row), max(1, max_col)
    try:
        raw = range_text.strip()
        if "!" in raw:
            raw = raw.split("!", 1)[1].strip()

        if ":" in raw:
            left, right = raw.split(":", 1)
            c1 = _col_number(left)
            r1 = _row_number(left) or 1
            c2 = _col_number(right)
            r2 = _row_number(right) or max_row
        else:
            # 单单元格或单个标识
            c1 = _col_number(raw)
            r1 = _row_number(raw) or 1
            c2 = c1
            r2 = r1

        start_row = min(r1, r2)
        end_row = max(r1, r2)
        start_col = min(c1, c2)
        end_col = max(c1, c2)

        return (
            max(1, start_row),
            max(1, start_col),
            min(max(1, max_row), end_row),
            min(max(1, max_col), end_col),
        )
    except (TypeError, ValueError):
        return 1, 1, max(1, max_row), max(1, max_col)


def _safe_stream_sheet(ws, headers_row: int, start_row: int, start_col: int, end_row: int, end_col: int, max_rows: int):
    """顺序遍历行，避免 ws.cell 随机寻址开销，并自适应探测有效行数。"""
    headers = []
    sample = []
    last_non_empty_row = 0

    # 使用 values_only 顺序读取
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        # 只要行内有非空内容，更新 last_non_empty_row
        if any(c is not None and str(c).strip() != "" for c in row):
            last_non_empty_row = row_idx

        if row_idx == headers_row and start_row <= headers_row:
            cleaned_row = [_cell(c) for c in row]
            headers = cleaned_row[start_col - 1 : end_col]
        elif headers_row < row_idx <= min(end_row, headers_row + max_rows):
            if row_idx >= start_row:
                cleaned_row = [_cell(c) for c in row]
                sliced = cleaned_row[start_col - 1 : end_col]
                if any(v is not None and str(v) != "" for v in sliced):
                    sample.append(sliced)

    return headers, sample, last_non_empty_row


def read_excel(
    path: Path,
    sheet: str | None = None,
    range_text: str | None = None,
    headers_row: int = 1,
    max_rows: int = 5,
) -> dict:
    """读取 Excel：无 sheet 时返回工作表清单；有 sheet 时返回表头与示例行。"""
    max_rows = max(1, min(int(max_rows), 100))
    headers_row = max(1, int(headers_row or 1))
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is None:
            sheets_meta = []
            for name in workbook.sheetnames:
                ws = workbook[name]
                # 跳过没有 iter_rows 的 chartsheet
                if not hasattr(ws, "iter_rows"):
                    continue
                sheets_meta.append({
                    "name": name,
                    "row_count": ws.max_row or 0,
                })
            return {
                "ok": True,
                "sheets": sheets_meta,
                "truncated": False,
            }

        if sheet not in workbook.sheetnames:
            return {"ok": False, "error": f"找不到工作表：{sheet}", "truncated": False}

        ws = workbook[sheet]
        if not hasattr(ws, "iter_rows"):
            return {"ok": False, "error": f"工作表 {sheet} 不是普通数据表（可能为图表页）", "truncated": False}

        max_row = ws.max_row or 0
        max_col = ws.max_column or 0

        # 如果 openpyxl 未能获取到有效的 max_row / max_col（常见于缺失 dimension 标签）
        effective_max_row = max_row if max_row > 0 else 1048576
        effective_max_col = max_col if max_col > 0 else 100

        start_row, start_col, end_row, end_col = _parse_range(
            range_text or "A1:ZZ1048576", effective_max_row, effective_max_col
        )

        headers, sample, last_row = _safe_stream_sheet(
            ws, headers_row, start_row, start_col, end_row, end_col, max_rows
        )

        actual_max_row = max(max_row, last_row)
        data_rows = max(0, actual_max_row - headers_row)

        return {
            "ok": True,
            "sheet": sheet,
            "headers": headers,
            "sample_rows": sample,
            "row_count": data_rows,
            "truncated": len(sample) < data_rows,
        }
    finally:
        workbook.close()


def search_excel(path: Path, keyword: str, max_hits: int = 6) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    hits = []
    try:
        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            if not hasattr(ws, "iter_rows"):
                continue
            for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                for col_idx, value in enumerate(row, start=1):
                    if value is None:
                        continue
                    if keyword.lower() in str(value).lower():
                        hits.append(
                            {
                                "sheet": sheet_name,
                                "row": row_idx,
                                "column": col_idx,
                                "value": str(value)[:80],
                            }
                        )
                        if len(hits) >= max_hits:
                            return {"ok": True, "keyword": keyword, "hits": hits, "truncated": True}
        return {"ok": True, "keyword": keyword, "hits": hits, "truncated": False}
    finally:
        workbook.close()
