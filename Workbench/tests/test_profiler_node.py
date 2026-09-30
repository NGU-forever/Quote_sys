"""Profiler 对异构资料的轻量来源定位测试。"""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from app.agent.readers.manifest import build_manifest
from app.agent.graph.state import AgentGraphState
from app.agent.graph.nodes.profiler import profiler_node
from app.core.config import settings


def _state(manifest: dict) -> AgentGraphState:
    return {
        "run_id": 1,
        "mode": "dry-run",
        "manifest": manifest,
        "db_stats": {},
        "sources_plan": {},
        "products_extracted": [],
        "packages_extracted": [],
        "bindings_extracted": [],
        "issues": [],
        "candidates": None,
        "validation_errors": [],
        "retry_count": 0,
        "status": "running",
    }


class ProfilerNodeTests(unittest.TestCase):
    def test_detects_excel_product_package_and_binding_sources(self):
        root = Path(__file__).resolve().parents[2] / "assets"
        with patch.object(settings, "assets_dir", str(root)):
            result = profiler_node(_state(build_manifest(root)))

        plan = result["sources_plan"]
        self.assertTrue(
            any(item["locator"].get("sheet") == "SKU主表" for item in plan["products"])
        )
        self.assertTrue(
            any(item["locator"].get("sheet") == "套餐定义表" for item in plan["packages"])
        )
        self.assertTrue(
            any(
                item["locator"].get("sheet") == "SKU主表"
                and item["locator"].get("column") == "适用设备"
                for item in plan["bindings"]
            )
        )

    def test_detects_markdown_sections_and_plain_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "notes.md").write_text(
                "# 商品清单\n"
                "| sku_id | category | name | min_price | unit | 适用设备 |\n"
                "| --- | --- | --- | --- | --- | --- |\n"
                "| P001 | device | 台式超声仪 | 2500 | 台 | C001 |\n"
                "\n## 套餐说明\n"
                "买2箱耗材赠1台台式超声仪。\n",
                encoding="utf-8",
            )
            (root / "plain.txt").write_text(
                "设备 P101 便携治疗仪，最低价 3000 元，设备。\n"
                "耗材 C101 电极片，最低价 2 元/袋，适用设备 P101。\n",
                encoding="utf-8",
            )
            with patch.object(settings, "assets_dir", tmp):
                result = profiler_node(_state(build_manifest(root)))

        plan = result["sources_plan"]
        self.assertTrue(any(item["file"] == "notes.md" for item in plan["products"]))
        self.assertTrue(
            any(
                item["file"] == "notes.md"
                and item["locator"].get("section") == "套餐说明"
                for item in plan["packages"]
            )
        )
        self.assertTrue(any(item["file"] == "plain.txt" for item in plan["products"]))

    def test_word_detects_table_and_promotion_paragraph(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            doc = Document()
            doc.add_heading("商品清单", level=1)
            table = doc.add_table(rows=2, cols=4)
            table.rows[0].cells[0].text = "sku_id"
            table.rows[0].cells[1].text = "category"
            table.rows[0].cells[2].text = "name"
            table.rows[0].cells[3].text = "最低价"
            table.rows[1].cells[0].text = "P001"
            table.rows[1].cells[1].text = "device"
            table.rows[1].cells[2].text = "台式超声仪"
            table.rows[1].cells[3].text = "2500"
            doc.add_heading("套餐说明", level=1)
            doc.add_paragraph("买2箱耗材赠1台台式超声仪。")
            doc.save(root / "business.docx")

            with patch.object(settings, "assets_dir", tmp):
                result = profiler_node(_state(build_manifest(root)))

        plan = result["sources_plan"]
        self.assertTrue(any(item["file"] == "business.docx" for item in plan["products"]))
        self.assertTrue(
            any(
                item["file"] == "business.docx"
                and item["locator"].get("section") == "套餐说明"
                for item in plan["packages"]
            )
        )

    def test_no_package_evidence_produces_empty_package_sources(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "products.txt").write_text(
                "设备 P001 超声波治疗仪，最低价 2500 元/台。\n"
                "耗材 C001 超声贴片，最低价 2 元/袋，适用设备 P001。\n",
                encoding="utf-8",
            )
            with patch.object(settings, "assets_dir", tmp):
                result = profiler_node(_state(build_manifest(root)))

        self.assertEqual(result["sources_plan"]["packages"], [])


if __name__ == "__main__":
    unittest.main()
