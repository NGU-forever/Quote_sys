"""LangGraph 多形态资料与套餐防发散端到端验收。"""
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from tests import db_support  # noqa: E402  # 统一测试库连接，必须在 app.* 之前导入

from docx import Document
from sqlalchemy import func, select
from sqlalchemy.orm import sessionmaker

from app.agent.graph.builder import build_import_graph
from app.agent.graph.runtime import GraphRuntime
from app.agent.llm_client import LLMResponse, ToolCall
from app.core.config import settings
from app.models.device_binding import DeviceBinding
from app.models.package import Package
from app.models.product import Product


class RecordingProvider:
    def __init__(self, payloads: dict):
        self.payloads = payloads
        self.calls: list[str] = []
        self._lock = threading.Lock()

    def chat(self, messages, tool_schemas):
        name = tool_schemas[0]["function"]["name"]
        with self._lock:
            self.calls.append(name)
        return LLMResponse(
            tool_calls=[ToolCall(name, self.payloads[name], call_id=name)],
            summary=name,
        )


def _db_session():
    engine = db_support.create_test_engine()
    db_support.reset_schema(engine)
    return sessionmaker(bind=engine, autoflush=False, autocommit=False)()


def _close_db(db) -> None:
    engine = db.get_bind()
    db.close()
    engine.dispose()


def _initial_state(mode: str = "dry-run") -> dict:
    return {
        "run_id": 1,
        "mode": mode,
        "manifest": {},
        "db_stats": {},
        "sources_plan": {},
        "products_extracted": [],
        "packages_extracted": [],
        "bindings_extracted": [],
        "issues": [],
        "candidates": None,
        "validation_errors": [],
        "retry_count": 0,
        "status": "running",
        "apply_result": {},
        "report": {},
    }


def _payloads() -> dict:
    return {
        "extract_products": {
            "products": [
                {
                    "sku_id": "P101",
                    "category": "device",
                    "name": "便携治疗仪",
                    "min_price": 3000,
                    "unit": "台",
                },
                {
                    "sku_id": "C101",
                    "category": "consumable",
                    "name": "治疗电极片",
                    "min_price": 2,
                    "unit": "袋",
                    "units_per_box": 100,
                },
            ],
            "issues": [],
        },
        "extract_packages": {
            "packages": [
                {
                    "package_id": "PKG101",
                    "name": "便携带耗材推广套餐",
                    "sku_ids": ["P101", "C101"],
                    "quantities": [1, 200],
                    "device_qty": 1,
                    "consumable_qty": 200,
                    "consumable_boxes": 2,
                    "fixed_bundle": False,
                    "note": "买2箱耗材赠1台便携治疗仪。",
                }
            ],
            "issues": [],
        },
        "extract_bindings": {
            "bindings": [
                {"device_sku_id": "P101", "consumable_sku_id": "C101"}
            ],
            "issues": [],
        },
    }


class AgentLangGraphE2ETests(unittest.TestCase):
    def test_word_business_document_apply_imports_all_entities(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            document = Document()
            document.add_heading("商品清单", level=1)
            table = document.add_table(rows=3, cols=5)
            table.rows[0].cells[0].text = "sku_id"
            table.rows[0].cells[1].text = "category"
            table.rows[0].cells[2].text = "name"
            table.rows[0].cells[3].text = "最低价"
            table.rows[0].cells[4].text = "适用设备"
            table.rows[1].cells[0].text = "P101"
            table.rows[1].cells[1].text = "device"
            table.rows[1].cells[2].text = "便携治疗仪"
            table.rows[1].cells[3].text = "3000"
            table.rows[1].cells[4].text = "P101"
            table.rows[2].cells[0].text = "C101"
            table.rows[2].cells[1].text = "consumable"
            table.rows[2].cells[2].text = "治疗电极片"
            table.rows[2].cells[3].text = "2"
            table.rows[2].cells[4].text = "P101"
            document.add_heading("套餐说明", level=1)
            document.add_paragraph("买2箱治疗电极片赠1台便携治疗仪。")
            document.save(root / "business.docx")

            provider = RecordingProvider(_payloads())
            db = _db_session()
            runtime = GraphRuntime(
                provider=provider,
                provider_name="mock",
                db=db,
                cancel_check=lambda: False,
            )
            graph = build_import_graph(runtime, None)
            try:
                with patch.object(settings, "assets_dir", tmp):
                    result = graph.invoke(_initial_state("apply"))
                counts = (
                    db.scalar(select(func.count()).select_from(Product)),
                    db.scalar(select(func.count()).select_from(Package)),
                    db.scalar(select(func.count()).select_from(DeviceBinding)),
                )
            finally:
                _close_db(db)

        self.assertEqual(result["status"], "succeeded")
        self.assertEqual(len(result["candidates"]["packages"]), 1)
        self.assertEqual(counts, (2, 1, 1))

    def test_plain_text_note_extracts_device_consumable_and_package(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "note.txt").write_text(
                "设备 P101 便携治疗仪，最低价 3000 元/台。\n"
                "耗材 C101 治疗电极片，最低价 2 元/袋，适用设备 P101。\n"
                "买2箱治疗电极片赠1台便携治疗仪。\n",
                encoding="utf-8",
            )
            provider = RecordingProvider(_payloads())
            db = _db_session()
            runtime = GraphRuntime(
                provider=provider,
                provider_name="mock",
                db=db,
                cancel_check=lambda: False,
            )
            graph = build_import_graph(runtime, None)
            try:
                with patch.object(settings, "assets_dir", tmp):
                    result = graph.invoke(_initial_state("dry-run"))
            finally:
                _close_db(db)

        self.assertEqual(result["status"], "succeeded")
        self.assertEqual(len(result["candidates"]["products"]), 2)
        self.assertEqual(len(result["candidates"]["packages"]), 1)
        self.assertEqual(len(result["candidates"]["bindings"]), 1)

    def test_package_node_is_skipped_when_source_has_no_package_evidence(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "prices.txt").write_text(
                "设备 P101 便携治疗仪，最低价 3000 元/台。\n"
                "耗材 C101 治疗电极片，最低价 2 元/袋，适用设备 P101。\n",
                encoding="utf-8",
            )
            payloads = _payloads()
            provider = RecordingProvider(payloads)
            db = _db_session()
            runtime = GraphRuntime(
                provider=provider,
                provider_name="mock",
                db=db,
                cancel_check=lambda: False,
            )
            graph = build_import_graph(runtime, None)
            try:
                with patch.object(settings, "assets_dir", tmp):
                    result = graph.invoke(_initial_state("dry-run"))
            finally:
                _close_db(db)

        self.assertEqual(result["candidates"]["packages"], [])
        self.assertNotIn("extract_packages", provider.calls)
        self.assertEqual(result["sources_plan"]["packages"], [])


if __name__ == "__main__":
    unittest.main()
